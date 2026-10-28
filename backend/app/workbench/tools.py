"""Executable capability registry, shared by discovery, validation and dispatch."""
from __future__ import annotations

import asyncio
import base64
import mimetypes
from pathlib import Path

import numpy as np

from ..agent.tools import search_payload
from ..design.brief import merge_facts
from ..design.analysis import analyze_design
from ..design.service import inventory_snapshot
from ..design.solver import solve_design
from ..inventory.evaluator import evaluate_candidates
from ..knowledge.ingestion import extract_text_from_file
from ..literature.service import run_literature_search
from ..models import AgentChatMessage, AgentTask, AgentToolCall, Attachment, GeneratedContent
from ..observability.usage import apply_usage_to_generated
from ..physics.laser_metrics import fit_power_curve
from .optics import OpticalScan, OpticalScan2D, OpticalSystem, analyze_system, scan_system, scan_system_2d
from .projects import SaveProject, ReadProject, read_project, save_project, project_from_design
from ..physics.caustic import fit_caustic
from .contracts import CausticData, ImportDesignProject
from ..physics.toolkit import run_coating_tmm, run_phase_match
from .contracts import (ArtifactQuery, AttachmentQuery, Coating, ComponentMatch, Finish, HistoryQuery,
                        InventoryQuery, KnowledgeQuery, LinearDesign, LiteratureQuery, PhaseMatch,
                        PowerCurve, SavedAnalysis, ToolHelp, ToolResultQuery)

REGISTRY = {
    "read_optical_project": (ReadProject, "Read latest or a prior version of the case optical engineering model, including owned parts, locked parameters, unknowns, and deterministic analysis."),
    "save_optical_project": (SaveProject, "Create/update the actual editable optical project and run its analysis. Use base_revision_id from latest read, preserve owned/locked parts and evidence. Every save is a new revision; never overwrites history."),
    "import_design_project": (ImportDesignProject, "Convert a saved linear design candidate into the editable optical project; preserve input/output mirror roles, crystal segments and owned-part locks. Reads actual saved numerical design, not a rewritten model."),
    "scan_optical_system_2d": (OpticalScan2D, "Compute a two-parameter T/S cavity stability map; supports physical gaps, radii, focal power and indices. Finite grid, no output-power optimization."),
    "fit_beam_caustic": (CausticData, "Fit measured air caustic radii vs z to find waist position/radius, divergence and equivalent M². Require actual data source and radius convention; never invent measurements."),
    "analyze_saved_design": (SavedAnalysis, "Precisely re-evaluate a saved specialized laser design: cavity/pump q traces, analytic waists and physical boundaries. Use this to investigate narrow pump waists instead of inferring a minimum from sampled artifact rows. Thermal focal length here is INSIDE the crystal; convert f_air=f_inside/n for the general system tool."),
    "read_tool_result": (ToolResultQuery, "Read a real tool output from an earlier run in this case, optionally a JSON key/index path. Reuse previous simulations and parameters without inventing their values."),
    "list_inventory": (InventoryQuery, "Search actual owned inventory, coating records and available quantities; paginate using offset. Empty inventory does not prevent proposing purchase specifications."),
    "read_attachment": (AttachmentQuery, "Read an attachment belonging to this case. Images are shown to the model; text/PDF/CSV can be paged. Read labels literally; identify uncertainty."),
    "read_artifact": (ArtifactQuery, "Read a previous saved artifact by id, optionally selecting a JSON key/index path. Generated designs are proposals, not measured facts."),
    "read_history": (HistoryQuery, "Read earlier task conversation including replies to questions; optional before_id for pagination."),
    "search_knowledge": (KnowledgeQuery, "Search accessible indexed documents and similar cases, returning evidence and citation identifiers."),
    "search_literature": (LiteratureQuery, "Search OpenAlex/arXiv metadata and abstracts, not full texts. Return actual records; do not invent literature when a source fails."),
    "analyze_optical_system": (OpticalSystem, "Compute general ordered ABCD system: single-pass beam propagation, standing-wave cavity or coplanar ring. T/S eigenmodes, stability, beam profiles, physical positions. Python engine; no gain/output-power solver."),
    "scan_optical_system": (OpticalScan, "Scan one specified element parameter in a general optical system; inspect stability and input/output beam size in both planes. Finite grid, no hidden optimization objective."),
    "fit_power_curve": (PowerCurve, "Fit measured pump/output data for threshold and slope efficiency with fit-quality warnings. Distinguish incident from absorbed pump power. Never fabricate measured points."),
    "compute_phase_match": (PhaseMatch, "Solve room-temperature SHG/SFG phase-matching and walk-off for BBO/LBO/KTP/BiBO. Not a conversion efficiency or thermal field simulation."),
    "match_components": (ComponentMatch, "Evaluate actual inventory against coating/ROC/aperture/quantity requirements; returns matched, uncertain and rejected candidates with reasons. Inspect_tool gives exact configuration."),
    "compute_coating": (Coating, "Evaluate thin-film stacks, measured curves, or clearly labeled nominal coating archetypes. inspect_tool gives configuration; nominal labels do not establish real off-design reflectivity."),
    "design_linear_laser": (LinearDesign, "Optional two-mirror end-pumped linear design: either mirror may be flat/concave; M1=input, M2=output. Extracts new owned-part constraints, preserves actual hardware, scans unfixed geometry. Saves detailed design; then import_design_project to create editable engineering model."),
    "inspect_tool": (ToolHelp, "Read supported configuration and model boundaries of a named tool before using unfamiliar fields."),
    "finish": (Finish, "Deliver results, evidence, assumptions, actionable questions and next steps. This ends the run, preserving all actual tool outputs."),
}


def schemas():
    result = []
    for name, (model, description) in REGISTRY.items():
        schema = model.model_json_schema()
        schema["properties"]["purpose"] = {"type": "string", "description": "Short user-visible reason for this operation"}
        result.append({"name": name, "description": description, "parameters": schema})
    return result


def compact(value, max_items=24, depth=0):
    """Explicit sampling for model context only; persisted tool output stays intact."""
    if isinstance(value, dict):
        return {key: compact(item, max_items, depth + 1) for key, item in value.items() if not key.startswith("_")}
    if isinstance(value, list):
        if len(value) > max_items:
            selected = set(int(x) for x in np.linspace(0, len(value) - 1, max_items))
            # Preserve narrow waists/extrema that even sampling can entirely miss.
            critical = set()
            for key in ("w_mm", "radius_mm", "m", "w_um"):
                if not all(isinstance(row, dict) and isinstance(row.get(key), (int, float)) for row in value):
                    continue
                ys = [row[key] for row in value]
                critical.update((min(range(len(ys)), key=ys.__getitem__), max(range(len(ys)), key=ys.__getitem__)))
                critical.update(i for i in range(1, len(ys) - 1) if (ys[i] < ys[i - 1] and ys[i] <= ys[i + 1]) or (ys[i] > ys[i - 1] and ys[i] >= ys[i + 1]))
            selected.update(sorted(critical)[:96])
            indexes = sorted(selected)
            return {"sampled_items": [compact(value[i], max_items, depth + 1) for i in indexes],
                    "total_items": len(value), "sampled_indices": indexes,
                    "all_detected_extrema_retained": len(critical) <= 96,
                    "note": "Sampled context with up to 96 numerical extrema retained. Raw curve may itself be coarse; never infer an exact physical waist from sampling alone. Use analytic waist output or analyze_saved_design."}
        return [compact(item, max_items, depth + 1) for item in value]
    if isinstance(value, str) and len(value) > 18000:
        return value[:18000] + "\n[Context text truncated; read the source in pages.]"
    return value


def image_payload(attachment):
    path = Path(attachment.filepath)
    mime = mimetypes.guess_type(attachment.filename)[0]
    if mime not in {"image/jpeg", "image/png", "image/webp", "image/gif"}:
        return None
    if path.stat().st_size > 12 * 1024 * 1024:
        raise ValueError("图片超过 12 MB，请压缩后重传。")
    data = path.read_bytes()
    return {"id": attachment.id, "data_url": f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"}


async def execute(name, args, *, db, case, task, provider, session_id, attachment_ids):
    if name == "read_optical_project":
        return {"status": "completed", "project_revision": read_project(db, case.id, args.revision_id)}
    if name == "save_optical_project":
        db.refresh(task)
        if task.status == "cancelled":
            raise ValueError("Task stopped before project save")
        allow = bool(args.unlock_quote.strip()) and args.unlock_quote.strip() in task.goal
        if args.unlock_quote and not allow:
            raise ValueError("Unlock quote must be present in the current user request")
        return save_project(db, case.id, args, allow_unlock=allow)
    if name == "import_design_project":
        artifact = db.get(GeneratedContent, args.content_id)
        if not artifact or artifact.case_id != case.id:
            raise ValueError("Design is outside this case")
        design = artifact.content.get("engineering_design", {})
        candidates = design.get("candidates", [])
        if not 1 <= args.candidate <= len(candidates):
            raise ValueError("Saved design has no such candidate")
        candidate = candidates[args.candidate-1]
        db.refresh(task)
        if task.status == "cancelled":
            raise ValueError("Task stopped before project import")
        previous = read_project(db, case.id)
        return save_project(db, case.id, SaveProject(base_revision_id=args.base_revision_id, project=project_from_design(design, candidate, previous), change_note=f"从设计 #{args.content_id} 候选 {args.candidate} 导入"))
    if name == "scan_optical_system_2d":
        return await asyncio.to_thread(scan_system_2d, args)
    if name == "fit_beam_caustic":
        return {"status": "completed", "source": args.data_source, **fit_caustic(args.z_mm, args.radius_mm, args.wavelength_nm)}
    if name == "analyze_saved_design":
        artifact = db.get(GeneratedContent, args.content_id)
        if not artifact or artifact.case_id != case.id:
            raise ValueError("Design artifact does not belong to this case")
        design = artifact.content.get("engineering_design", {})
        candidates = design.get("candidates", [])
        if not 1 <= args.candidate <= len(candidates):
            raise ValueError("Saved design has no such candidate")
        result = await asyncio.to_thread(analyze_design, design, candidates[args.candidate - 1], scan_samples=150,
                                         **args.model_dump(exclude={"content_id", "candidate"}))
        return {"status": "completed", "content_id": artifact.id, "candidate": args.candidate, "analysis": result}
    if name == "read_tool_result":
        call = db.query(AgentToolCall).join(AgentTask).filter(AgentToolCall.id == args.tool_call_id, AgentTask.case_id == case.id).first()
        if not call or call.tool_name == "workbench_request":
            raise ValueError("Saved tool result does not belong to this case")
        output = call.output_json
        for key in args.path.split(".") if args.path else []:
            output = output[int(key)] if isinstance(output, list) else output[key]
        return {"status": "completed", "source_tool_call_id": call.id, "source_status": call.status, "name": call.tool_name,
                "original_input": call.input_json, "output": output}
    if name == "analyze_optical_system":
        return await asyncio.to_thread(analyze_system, args)
    if name == "scan_optical_system":
        return await asyncio.to_thread(scan_system, args)
    if name == "fit_power_curve":
        if len(args.pump_w) != len(args.output_w) or min(args.pump_w + args.output_w) < 0 or len(set(args.pump_w)) < 3:
            raise ValueError("功率数组必须等长、非负且至少包含三个不同泵浦功率。")
        result = await asyncio.to_thread(fit_power_curve, args.pump_w, args.output_w)
        return {**result, "pump_basis": args.pump_basis, "data_source": args.data_source,
                "data": {"pump_w": args.pump_w, "output_w": args.output_w}}
    if name == "compute_phase_match":
        return await asyncio.to_thread(run_phase_match, args.model_dump(exclude_none=True))
    if name == "compute_coating":
        return await asyncio.to_thread(run_coating_tmm, args.config)
    if name == "list_inventory":
        items = inventory_snapshot(db)
        filtered = [item for item in items if (not args.category or item["category"] == args.category)
                    and (not args.query or args.query.lower() in f"{item['name']} {item['raw_spec']}".lower())]
        return {"status": "completed", "total": len(filtered), "offset": args.offset,
                "items": filtered[args.offset:args.offset + args.limit]}
    if name == "match_components":
        return evaluate_candidates(db, args.requirement)
    if name == "search_knowledge":
        return search_payload(db, args.query, case.id, args.top_k, task.id)
    if name == "search_literature":
        return await asyncio.to_thread(run_literature_search, args.model_dump(exclude_none=True))
    if name == "read_history":
        query = db.query(AgentChatMessage).filter_by(session_id=session_id)
        if args.before_id:
            query = query.filter(AgentChatMessage.id < args.before_id)
        rows = query.order_by(AgentChatMessage.id.desc()).limit(args.limit + 1).all()
        return {"messages": [{"id": row.id, "role": row.role, "content": row.content, "metadata": row.metadata_json}
                             for row in reversed(rows[:args.limit])], "has_older": len(rows) > args.limit}
    if name == "read_artifact":
        item = db.get(GeneratedContent, args.content_id)
        if not item or item.case_id != case.id:
            raise ValueError("Artifact does not belong to this case")
        value = item.content
        for key in args.path.split(".") if args.path else []:
            value = value[int(key)] if isinstance(value, list) else value[key]
        notes = []
        if "engineering_design" in item.content:
            notes = ["Legacy thermal f is a physical focal length INSIDE crystal; general optical-system air-equivalent focal length is f/n. Compare only after conversion.",
                     "Saved pump.points are a coarse trace; read trace.waists from analyze_saved_design for exact waist locations. Absence of a point in compacted output does not prove a displaced focus."]
        return {"content_id": item.id, "path": args.path, "content": value, "interpretation_notes": notes}
    if name == "read_attachment":
        item = db.get(Attachment, args.attachment_id)
        if not item or item.case_id != case.id:
            raise ValueError("Attachment does not belong to this case")
        picture = image_payload(item)
        if picture:
            return {"status": "completed", "attachment_id": item.id, "filename": item.filename,
                    "message": "Image supplied to model; reading is document evidence, not a measurement.", "_images": [picture]}
        if Path(item.filename).suffix.lower() in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            book = load_workbook(item.filepath, read_only=True, data_only=True)
            try:
                lines = []
                for sheet in book:
                    lines.append(f"Sheet: {sheet.title}; rows={sheet.max_row}")
                    for row in sheet.iter_rows(values_only=True):
                        lines.append("\t".join(str(v) if v is not None else "" for v in row))
                text = "\n".join(lines)
            finally:
                book.close()
        else:
            text = await asyncio.to_thread(extract_text_from_file, item.filepath, item.file_type)
        return {"status": "completed", "attachment_id": item.id, "filename": item.filename,
                "text": text[args.offset:args.offset + args.limit], "total_characters": len(text),
                "next_offset": args.offset + args.limit if args.offset + args.limit < len(text) else None}
    if name == "inspect_tool":
        if args.name not in REGISTRY:
            raise ValueError("Unknown tool; choose a name from the available registry")
        help_text = REGISTRY[args.name][1]
        if args.name == "compute_coating":
            help_text += "\nconfig: mode ('auto'|'exact_stack'|'vendor_curve'|'archetype'), query_wavelengths_nm (array), aoi_deg (default 0), polarization ('s'|'p'|'avg'). Exact stack: stack={incident:'air',substrate:'fused_silica',layers:[{thickness_nm:100,material:'sio2_film'}]} (layer n may replace material). Measured curve: curve={wl_nm:[...],R:[fractions 0..1],T:[fractions 0..1],source:'document id'}. Nominal archetype: nominal={function:'HR',design_wavelength_nm:1064,target_R:0.995}. Nominal results describe a representative stack, never the actual part. Measured curves are not extrapolated."
        if args.name == "match_components":
            help_text += "\nrequirement: category, quantity, roc_mm ('flat' or signed mm), roc_tol_pct, min_diameter_mm, max_results; surfaces: [{surface:'S1', wavelength_nm:1064, function:'HR', min_R_pct:99.8}]; phase_match: {crystal:'lbo',lambda1_nm:1064,pm_type:'I'}."
        return {"name": args.name, "description": help_text, "parameters": REGISTRY[args.name][0].model_json_schema()}
    if name == "design_linear_laser":
        inventory = inventory_snapshot(db)
        # Extract from actual human input, not the model-generated tool arguments.
        context = {"title": case.title, "description": case.description or "", "message": task.goal,
                   "existing_facts": (case.parameters or {}).get("design_facts", {}), "inventory": inventory[:150]}
        images = []
        for aid in attachment_ids:
            item = db.get(Attachment, aid)
            if item and item.case_id == case.id and image_payload(item):
                images.append({"id": aid, "mime": mimetypes.guess_type(item.filename)[0], "data": Path(item.filepath).read_bytes()})
        extraction = await provider.extract_design_facts(context, images)
        facts, changes, notes = merge_facts(context["existing_facts"], extraction,
                                            {key: context[key] for key in ("title", "description", "message")}, attachment_ids)
        design = await asyncio.to_thread(solve_design, facts, inventory)
        db.refresh(task)
        if task.status == "cancelled":
            raise ValueError("Task stopped; no new design or case facts were saved")
        design.update({"changes": changes, "intake_notes": notes})
        content = {"engineering_design": design, "summary": design.get("summary", ""), "_usage": extraction.get("_usage", {})}
        artifact = GeneratedContent(case_id=case.id, content_type="plan", content=content, prompt_version="workbench-linear-tool-v1")
        apply_usage_to_generated(artifact, content)
        db.add(artifact)
        db.flush()
        case.parameters = {**(case.parameters or {}), "design_facts": facts}
        return {"status": design["status"], "content_id": artifact.id, "engineering_design": design}
    raise ValueError("Unknown executable tool")
